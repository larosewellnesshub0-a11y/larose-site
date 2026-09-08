#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Build the clinic's CONTENT-MAP.xlsx from the live JSON content files.

The workbook gives a non-technical editor one calm, bilingual place to review
the site's copy.  Its final _ref column preserves the exact JSON location, so
the returned workbook can be applied safely by tools/apply_content_map.py
without asking the client to understand the content schema.

Run:  python tools/export_content_map.py
"""

import io
import json
import os
import re
import sys
from dataclasses import dataclass

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT_DIR = os.path.join(ROOT, "content")
OUTPUT = os.path.join(ROOT, "_project", "CONTENT-MAP.xlsx")

CONTENT_FILES = (
    "site.json",
    "pages.json",
    "specialties.json",
    "doctors.json",
    "branches.json",
    "articles.json",
    "digital.json",
    "reviews.json",
)

HEADERS = (
    "Page / Section",
    "Part",
    "Link",
    "Aim",
    "Current (AR)",
    "Current (EN)",
    "New (AR)",
    "New (EN)",
    "Notes",
    "_ref",
)

# Brand colours shared with tools/build_content_sheet.py.
OLIVE = "333524"
OLIVE_LIGHT = "E7E6CF"
CHAMPAGNE = "D4B793"
PAPER = "FBF9F4"
LINE = "E2DECE"
EDIT_FILL = "FFF0D8"

HEADER_FILL = PatternFill("solid", fgColor=OLIVE) if Workbook else None
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF") if Workbook else None
BODY_FONT = Font(name="Calibri", size=10, color=OLIVE) if Workbook else None
REF_FONT = Font(name="Calibri", size=8, color="777777") if Workbook else None
NEW_FILL = PatternFill("solid", fgColor=EDIT_FILL) if Workbook else None
ALT_FILL = PatternFill("solid", fgColor=PAPER) if Workbook else None
THIN = Side(style="thin", color=LINE) if Workbook else None
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN) if Workbook else None
WRAP = Alignment(wrap_text=True, vertical="top") if Workbook else None
WRAP_RTL = Alignment(wrap_text=True, vertical="top", readingOrder=2) if Workbook else None


# These fields control routing, publication, media, relationships or factual
# metadata.  They are not prose and must not enter a copy-editing round trip.
MACHINE_KEYS = {
    "slug", "id", "index", "published", "featured", "icon", "image",
    "imageplaceholder", "imagepath", "imagepaths", "src", "portrait",
    "portrait2x", "portraitplaceholder", "date", "dates", "checked",
    "createdat", "updatedat", "modifiedat", "datemodified", "openingdate",
    "timestamp", "geo", "lat", "lng", "coordinate", "coordinates",
    "showprice", "href", "url",
    "mapsurl", "domain", "tel", "number", "formsendpoint", "key", "type",
    "category", "specialty", "author", "reviewedby", "code", "dir",
    "locale", "childrenfrom", "linksfrom", "pluscode", "source",
    "status", "sample", "staffed", "enabled", "mega", "illustrative",
    "isprimary", "placeholder", "consent", "hasbeforeafter", "hasreviews",
    "readingtime", "count", "value", "rating", "ekshefclinicid",
    "latinterms", "n",
}

MACHINE_COLLECTIONS = {"specialties", "branches"}
IMAGE_SUFFIXES = (".avif", ".gif", ".jpeg", ".jpg", ".png", ".svg", ".webp")


def bilingual(ar, en):
    """One bilingual line, with the two reading directions clearly divided."""
    return f"{ar} / {en}"


def load_json(filename):
    path = os.path.join(CONTENT_DIR, filename)
    with io.open(path, "r", encoding="utf-8") as handle:
        return json.load(handle)


def scalar_text(value):
    """Excel cells need strings; None represents an existing empty field."""
    if value is None:
        return ""
    return str(value)


def pointer(tokens):
    """Return an RFC 6901 JSON Pointer for path tokens."""
    if not tokens:
        return ""
    escaped = (str(token).replace("~", "~0").replace("/", "~1") for token in tokens)
    return "/" + "/".join(escaped)


def single_ref(filename, tokens):
    return f"{filename}#{pointer(tokens)}"


def paired_ref(filename, ar_tokens, en_tokens):
    """Two exact pointers are needed for locale-separated parallel arrays."""
    return f"{single_ref(filename, ar_tokens)} || {single_ref(filename, en_tokens)}"


def is_scalar_copy(value):
    return value is None or isinstance(value, str)


def looks_like_machine_value(value):
    if not isinstance(value, str):
        return False
    stripped = value.strip()
    lowered = stripped.lower().split("?", 1)[0].split("#", 1)[0]
    if re.match(r"^(?:https?:|mailto:|tel:)", stripped, re.I):
        return True
    if lowered.startswith(("assets/", "/assets/")):
        return True
    if lowered.endswith(IMAGE_SUFFIXES):
        return True
    return False


def should_skip_key(filename, tokens, key, value):
    """Keep prose and exact _todo fields; discard operational data."""
    if key.startswith("_"):
        return key != "_todo"

    lowered = key.lower()
    # `reviewedBy` is a doctor-slug relationship in articles.json, but the
    # same spelling is visible interface copy under site.ui.
    if filename == "site.json" and tokens[:1] == ("ui",) and lowered == "reviewedby":
        return False
    if lowered in MACHINE_KEYS:
        return True

    # `specialties` and `branches` are top-level record collections in their
    # own files, but inside a record they are arrays of slug references.
    if lowered in MACHINE_COLLECTIONS and isinstance(value, list):
        if all(not isinstance(item, (dict, list)) for item in value):
            return True

    # Flags and counts are never editable copy, even when a future schema gives
    # one a new key that is not yet listed above.
    if isinstance(value, (bool, int, float)):
        return True

    # File paths and external links are machine values regardless of their key.
    if looks_like_machine_value(value):
        return True

    # Empty media slots do not look like paths yet, so their schema key has to
    # carry the decision. `before`/`after` remain valid visible UI labels when
    # their value is a bilingual object rather than a scalar media field.
    if lowered in ("before", "after") and is_scalar_copy(value):
        return True
    if tokens and tokens[-1] == "logo" and isinstance(value, str):
        return True

    # Contact numbers are operational facts, not language-specific prose.
    if tokens and tokens[-1] in ("phone", "whatsapp") and lowered == "display":
        return True
    if lowered == "whatsapp" and isinstance(value, str) and value.strip().isdigit():
        return True

    # This is a bidi/font allow-list used by the renderer, not published copy.
    if lowered == "latinterms":
        return True

    return False


def pair_shape(node):
    """Identify the two bilingual shapes used by the current content files."""
    if not isinstance(node, dict) or "ar" not in node or "en" not in node:
        return None
    ar_value, en_value = node["ar"], node["en"]
    if is_scalar_copy(ar_value) and is_scalar_copy(en_value):
        return "scalar"
    if isinstance(ar_value, list) and isinstance(en_value, list):
        if all(is_scalar_copy(item) for item in ar_value + en_value):
            return "parallel-lists"
    return None


def name_pair(record, fallback_ar, fallback_en):
    if not isinstance(record, dict):
        return fallback_ar, fallback_en
    for key in ("name", "title", "label", "shortName"):
        value = record.get(key)
        if isinstance(value, dict):
            ar = scalar_text(value.get("ar")).strip()
            en = scalar_text(value.get("en")).strip()
            if ar or en:
                return ar or fallback_ar, en or fallback_en
    return fallback_ar, fallback_en


def indexed_record(root, collection, tokens):
    """Return (record, index) when tokens point inside a record collection."""
    if len(tokens) < 2 or tokens[0] != collection or not isinstance(tokens[1], int):
        return None, 0
    records = root.get(collection, []) if isinstance(root, dict) else []
    index = tokens[1]
    if 0 <= index < len(records) and isinstance(records[index], dict):
        return records[index], index
    return None, index


def base_url(domain):
    domain = scalar_text(domain).strip().rstrip("/")
    if not domain:
        return ""
    if re.match(r"^https?://", domain, re.I):
        return domain
    return "https://" + domain


def live_url(domain, path):
    base = base_url(domain)
    if not base or not path:
        return ""
    return f"{base}/ar/{path.lstrip('/')}"


@dataclass(frozen=True)
class Location:
    tour: int
    entity: int
    heading: str
    link: str


STATIC_PAGES = {
    "home": (1, bilingual("الرئيسية", "Home"), "index.html"),
    "specialties": (2, bilingual("التخصصات", "Specialties"), "specialties/index.html"),
    "doctors": (3, bilingual("الأطباء", "Doctors"), "doctors/index.html"),
    "branches": (4, bilingual("الفروع", "Branches"), "branches/index.html"),
    "about": (5, bilingual("عن لاروز", "About La Rose"), "about/index.html"),
    "contact": (5, bilingual("اتصل بنا", "Contact"), "contact.html"),
    "patients": (5, bilingual("دليل المريض", "Patient guide"), "patients/index.html"),
    "booking": (5, bilingual("دليل المريض › حجز موعد", "Patient guide > Booking"), "patients/booking.html"),
    "first-visit": (5, bilingual("دليل المريض › أول زيارة", "Patient guide > First visit"), "patients/first-visit.html"),
    "preparation": (5, bilingual("دليل المريض › التحضير للكشف", "Patient guide > Visit preparation"), "patients/preparation.html"),
    "faq": (5, bilingual("دليل المريض › الأسئلة الشائعة", "Patient guide > FAQ"), "patients/faq.html"),
    "rights": (5, bilingual("دليل المريض › حقوق المريض", "Patient guide > Patient rights"), "patients/rights.html"),
    "home-visits": (5, bilingual("دليل المريض › الزيارات المنزلية", "Patient guide > Home visits"), "patients/home-visits.html"),
    "tools": (5, bilingual("أدوات طبية", "Medical tools"), "tools/index.html"),
    "legal": (5, bilingual("الصفحات القانونية", "Legal pages"), "legal/disclaimer.html"),
    "articles": (6, bilingual("المركز المعرفي", "Knowledge Centre"), "articles/index.html"),
    "digital": (7, bilingual("لاروز ديچيتال", "La Rose Digital"), "digital/index.html"),
    "reviews": (8, bilingual("آراء المرضى", "Patient reviews"), "about/reviews.html"),
}


def site_location(tokens, root, domain):
    first = tokens[0] if tokens else "site"
    if first == "proof":
        return Location(1, 0, bilingual("الرئيسية › أرقام لاروز", "Home > La Rose in numbers"),
                        live_url(domain, "index.html"))
    if first == "media":
        return Location(1, 1, bilingual("الرئيسية › الصورة الرئيسية", "Home > Hero image"),
                        live_url(domain, "index.html"))
    if first == "sharing":
        return Location(6, -1, bilingual("المركز المعرفي › المشاركة والذكاء الاصطناعي",
                                         "Knowledge Centre > Sharing and Ask AI"),
                        live_url(domain, "articles/index.html"))

    headings = {
        "brand": bilingual("إعدادات الموقع › الهوية", "Site-wide settings > Brand"),
        "contact": bilingual("إعدادات الموقع › التواصل", "Site-wide settings > Contact"),
        "hours": bilingual("إعدادات الموقع › مواعيد العمل", "Site-wide settings > Opening hours"),
        "nav": bilingual("إعدادات الموقع › القائمة الرئيسية", "Site-wide settings > Main navigation"),
        "footer": bilingual("إعدادات الموقع › الفوتر", "Site-wide settings > Footer"),
        "legal": bilingual("إعدادات الموقع › النصوص القانونية", "Site-wide settings > Legal copy"),
        "ui": bilingual("إعدادات الموقع › نصوص الواجهة", "Site-wide settings > Interface labels"),
        "social": bilingual("إعدادات الموقع › السوشيال", "Site-wide settings > Social links"),
    }
    heading = headings.get(first, bilingual("إعدادات الموقع", "Site-wide settings"))

    # Add the visible menu/footer item, where possible, without turning the
    # section heading into another JSON path.
    if first == "nav" and len(tokens) >= 2 and isinstance(tokens[1], int):
        nav = root.get("nav", [])
        if tokens[1] < len(nav):
            ar, en = name_pair(nav[tokens[1]], "عنصر", "Item")
            heading = bilingual(f"إعدادات الموقع › القائمة › {ar}",
                                f"Site-wide settings > Navigation > {en}")
    return Location(0, 0, heading, "")


def pages_location(tokens, root, domain):
    key = str(tokens[0]) if tokens else "pages"
    page = STATIC_PAGES.get(key)
    if page:
        tour, heading, path = page
        return Location(tour, 0, heading, live_url(domain, path))
    return Location(5, 99, bilingual("صفحات المرضى والمعلومات", "Patient and information pages"), "")


def specialty_location(tokens, root, domain):
    record, index = indexed_record(root, "specialties", tokens)
    ar, en = name_pair(record, f"تخصص {index + 1}", f"Specialty {index + 1}")
    heading = bilingual(f"التخصصات › {ar}", f"Specialties > {en}")
    slug = record.get("slug", "") if record else ""

    if record and "treatments" in tokens:
        position = tokens.index("treatments")
        if len(tokens) > position + 1 and isinstance(tokens[position + 1], int):
            treatment_index = tokens[position + 1]
            treatments = record.get("treatments", [])
            if treatment_index < len(treatments):
                tar, ten = name_pair(treatments[treatment_index], "خدمة", "Service")
                heading = bilingual(f"التخصصات › {ar} › {tar}",
                                    f"Specialties > {en} > {ten}")
    return Location(2, index, heading,
                    live_url(domain, f"specialties/{slug}.html") if slug else "")


def doctor_location(tokens, root, domain):
    record, index = indexed_record(root, "doctors", tokens)
    ar, en = name_pair(record, f"طبيب {index + 1}", f"Doctor {index + 1}")
    slug = record.get("slug", "") if record else ""
    return Location(3, index, bilingual(f"الأطباء › {ar}", f"Doctors > {en}"),
                    live_url(domain, f"doctors/{slug}.html") if slug else "")


def branch_location(tokens, root, domain):
    record, index = indexed_record(root, "branches", tokens)
    ar, en = name_pair(record, f"فرع {index + 1}", f"Branch {index + 1}")
    slug = record.get("slug", "") if record else ""
    return Location(4, index, bilingual(f"الفروع › {ar}", f"Branches > {en}"),
                    live_url(domain, f"branches/{slug}.html") if slug else "")


def article_location(tokens, root, domain):
    collection = tokens[0] if tokens else "articles"
    if collection == "categories":
        record, index = indexed_record(root, "categories", tokens)
        ar, en = name_pair(record, f"تصنيف {index + 1}", f"Category {index + 1}")
        slug = record.get("slug", "") if record else ""
        return Location(6, index, bilingual(f"المركز المعرفي › تصنيف {ar}",
                                             f"Knowledge Centre > {en} category"),
                        live_url(domain, f"articles/category-{slug}.html") if slug else "")

    record, index = indexed_record(root, "articles", tokens)
    ar, en = name_pair(record, f"محتوى {index + 1}", f"Entry {index + 1}")
    slug = record.get("slug", "") if record else ""
    return Location(6, 1000 + index, bilingual(f"المركز المعرفي › {ar}",
                                                f"Knowledge Centre > {en}"),
                    live_url(domain, f"articles/{slug}.html") if slug else "")


def digital_location(tokens, root, domain):
    if tokens and tokens[0] == "_internalPricing":
        return Location(7, 99, bilingual("لاروز ديچيتال › أسئلة داخلية معلّقة",
                                         "La Rose Digital > Outstanding internal questions"), "")
    record, index = indexed_record(root, "products", tokens)
    ar, en = name_pair(record, f"منتج {index + 1}", f"Product {index + 1}")
    slug = record.get("slug", "") if record else ""
    heading = bilingual(f"لاروز ديچيتال › {ar}", f"La Rose Digital > {en}")
    if record and "formats" in tokens:
        position = tokens.index("formats")
        if len(tokens) > position + 1 and isinstance(tokens[position + 1], int):
            format_index = tokens[position + 1]
            formats = record.get("formats", [])
            if format_index < len(formats):
                far, fen = name_pair(formats[format_index], "نظام متابعة", "Programme format")
                heading = bilingual(f"لاروز ديچيتال › {ar} › {far}",
                                    f"La Rose Digital > {en} > {fen}")
    return Location(7, index, heading,
                    live_url(domain, f"digital/{slug}.html") if slug else live_url(domain, "digital/index.html"))


def reviews_location(tokens, root, domain):
    first = tokens[0] if tokens else "reviews"
    if first == "beforeAfter":
        return Location(8, 2, bilingual("النتائج › قبل وبعد", "Results > Before and after"),
                        live_url(domain, "about/before-after.html"))
    if first == "topics":
        return Location(8, 0, bilingual("آراء المرضى › أكتر حاجات مذكورة",
                                         "Patient reviews > Most-mentioned topics"),
                        live_url(domain, "about/reviews.html"))
    if first == "pending":
        return Location(8, 3, bilingual("آراء المرضى › في انتظار المراجعة",
                                         "Patient reviews > Awaiting moderation"), "")
    return Location(8, 1, bilingual("آراء المرضى", "Patient reviews"),
                    live_url(domain, "about/reviews.html"))


def location_for(filename, tokens, root, domain):
    if filename == "site.json":
        return site_location(tokens, root, domain)
    if filename == "pages.json":
        return pages_location(tokens, root, domain)
    if filename == "specialties.json":
        return specialty_location(tokens, root, domain)
    if filename == "doctors.json":
        return doctor_location(tokens, root, domain)
    if filename == "branches.json":
        return branch_location(tokens, root, domain)
    if filename == "articles.json":
        return article_location(tokens, root, domain)
    if filename == "digital.json":
        return digital_location(tokens, root, domain)
    if filename == "reviews.json":
        return reviews_location(tokens, root, domain)
    return Location(99, 0, bilingual("محتوى آخر", "Other content"), "")


UI_PARTS = {
    "bookNow": bilingual("نص زر حجز الموعد", "Book-appointment button label"),
    "bookShort": bilingual("نص زر الحجز المختصر", "Short booking button label"),
    "whatsapp": bilingual("نص زر واتساب", "WhatsApp button label"),
    "whatsappShort": bilingual("نص واتساب المختصر", "Short WhatsApp label"),
    "callUs": bilingual("نص زر الاتصال", "Call button label"),
    "readMore": bilingual("نص رابط اقرأ المزيد", "Read-more link label"),
    "viewAll": bilingual("نص رابط عرض الكل", "View-all link label"),
    "learnMore": bilingual("نص رابط اعرف أكثر", "Learn-more link label"),
    "viewSpecialty": bilingual("نص رابط تفاصيل التخصص", "Specialty-details link label"),
    "viewProfile": bilingual("نص رابط ملف الطبيب", "Doctor-profile link label"),
    "skipToContent": bilingual("رابط تخطي القائمة", "Skip-to-content label"),
    "menu": bilingual("نص زر القائمة", "Menu button label"),
    "close": bilingual("نص زر الإغلاق", "Close button label"),
    "home": bilingual("اسم رابط الرئيسية", "Home link label"),
    "comingSoon": bilingual("رسالة قريباً", "Coming-soon message"),
    "sample": bilingual("شارة المثال التوضيحي", "Sample-content badge"),
    "openingSoon": bilingual("رسالة الفرع سيفتح قريباً", "Opening-soon message"),
    "openNow": bilingual("رسالة الفرع مفتوح", "Open-now message"),
    "tabOverview": bilingual("اسم تبويب النظرة العامة", "Overview tab label"),
    "tabDoctors": bilingual("اسم تبويب الأطباء", "Doctors tab label"),
    "tabReviews": bilingual("اسم تبويب الآراء", "Reviews tab label"),
    "tabResults": bilingual("اسم تبويب النتائج", "Results tab label"),
    "tabFaq": bilingual("اسم تبويب الأسئلة", "FAQ tab label"),
    "conditionsTreated": bilingual("عنوان الحالات التي نعالجها", "Conditions-treated heading"),
    "treatments": bilingual("عنوان الخدمات والعلاجات", "Treatments heading"),
    "ourDoctors": bilingual("عنوان أطباء القسم", "Department-doctors heading"),
    "clinicDays": bilingual("عنوان أيام العيادة", "Clinic-days label"),
    "addReview": bilingual("نص إضافة تجربة", "Add-review label"),
    "before": bilingual("وسم قبل", "Before label"),
    "after": bilingual("وسم بعد", "After label"),
    "selectSpecialty": bilingual("طلب اختيار التخصص", "Specialty selector prompt"),
    "selectDoctor": bilingual("طلب اختيار الطبيب", "Doctor selector prompt"),
    "selectBranch": bilingual("طلب اختيار الفرع", "Branch selector prompt"),
    "findAppointment": bilingual("نص البحث عن موعد", "Find-appointment label"),
    "priceOnConsult": bilingual("رسالة السعر بعد الكشف", "Price-after-consultation message"),
    "homeVisits": bilingual("اسم الزيارات المنزلية", "Home-visits label"),
    "onlineFollowUp": bilingual("اسم المتابعة أونلاين", "Online-follow-up label"),
    "healthTools": bilingual("اسم الأدوات الطبية", "Health-tools label"),
    "knowledgeCentre": bilingual("اسم المركز المعرفي", "Knowledge Centre label"),
    "articles": bilingual("اسم المقالات", "Articles label"),
    "updates": bilingual("اسم المستجدات", "Updates label"),
    "qa": bilingual("اسم اسأل الطبيب", "Ask-the-doctor label"),
    "tips": bilingual("اسم النصائح اليومية", "Daily-tips label"),
    "readingTime": bilingual("وحدة وقت القراءة", "Reading-time unit"),
    "publishedOn": bilingual("وسم تاريخ النشر", "Published-date label"),
    "reviewedBy": bilingual("وسم المراجعة الطبية", "Medical-review label"),
    "toolDisclaimer": bilingual("تنبيه الأدوات الطبية", "Health-tool notice"),
    "calculate": bilingual("نص زر احسب", "Calculate button label"),
    "reset": bilingual("نص زر إعادة", "Reset button label"),
    "saveLocal": bilingual("تنبيه حفظ البيانات محلياً", "Local-storage notice"),
}

GENERIC_PARTS = {
    "shortName": bilingual("الاسم المختصر", "Short name"),
    "kind": bilingual("وصف نوع العيادة", "Clinic-type description"),
    "tagline": bilingual("الجملة التعريفية", "Tagline"),
    "line1": bilingual("السطر الأول في الشعار", "Wordmark line 1"),
    "line2": bilingual("السطر الثاني في الشعار", "Wordmark line 2"),
    "label": bilingual("النص الظاهر", "Visible label"),
    "desc": bilingual("الوصف المختصر", "Short description"),
    "description": bilingual("الوصف", "Description"),
    "sub": bilingual("العنوان المساند", "Supporting headline"),
    "intro": bilingual("المقدمة", "Introduction"),
    "lede": bilingual("الافتتاحية", "Opening paragraph"),
    "short": bilingual("الوصف المختصر", "Short description"),
    "summary": bilingual("الملخص", "Summary"),
    "heading": bilingual("عنوان القسم", "Section heading"),
    "body": bilingual("نص القسم", "Section copy"),
    "q": bilingual("سؤال شائع", "FAQ question"),
    "a": bilingual("إجابة السؤال", "FAQ answer"),
    "alt": bilingual("النص البديل للصورة", "Image alt text"),
    "caption": bilingual("تعليق الصورة", "Image caption"),
    "bio": bilingual("نبذة الطبيب", "Doctor biography"),
    "days": bilingual("أيام العيادة", "Clinic days"),
    "hours": bilingual("ساعات العمل", "Opening hours"),
    "display": bilingual("النص المعروض", "Displayed text"),
    "bookingNote": bilingual("ملاحظة الحجز", "Booking note"),
    "address": bilingual("العنوان", "Address"),
    "landmark": bilingual("علامة مميزة", "Landmark"),
    "area": bilingual("المنطقة", "Area"),
    "city": bilingual("المدينة", "City"),
    "country": bilingual("الدولة", "Country"),
    "blurb": bilingual("نبذة الفوتر", "Footer introduction"),
    "disclaimer": bilingual("إخلاء المسؤولية الطبية", "Medical disclaimer"),
    "resultsVary": bilingual("تنبيه اختلاف النتائج", "Results-vary notice"),
    "sampleContent": bilingual("تنبيه المحتوى التوضيحي", "Sample-content notice"),
    "copyright": bilingual("حقوق النشر", "Copyright line"),
    "cta": bilingual("نص الدعوة للتصرف", "Call-to-action label"),
    "excerpt": bilingual("مقتطف المحتوى", "Entry excerpt"),
    "question": bilingual("السؤال الجاهز للذكاء الاصطناعي", "AI question prompt"),
    "note": bilingual("ملاحظة للقارئ", "Reader note"),
    "copied": bilingual("رسالة نسخ الرابط", "Copy-link confirmation"),
    "text": bilingual("النص", "Copy"),
}

LIST_PARTS = {
    "credentials": bilingual("مؤهل أو خبرة", "Credential or experience point"),
    "treats": bilingual("حالة نعالجها", "Condition treated"),
    "facts": bilingual("معلومة سريعة عن الخدمة", "Service fact"),
    "gettingHere": bilingual("خطوة للوصول للفرع", "Directions step"),
    "inside": bilingual("ميزة داخل المنتج", "What's-inside point"),
    "audience": bilingual("حالة يناسبها المنتج", "Intended-audience point"),
    "includes": bilingual("بند مشمول", "Included item"),
    "tags": bilingual("وسم موضوع", "Topic tag"),
}


def semantic_key(tokens):
    if not tokens:
        return "copy"
    if isinstance(tokens[-1], int) and len(tokens) >= 2:
        return str(tokens[-2])
    return str(tokens[-1])


def humanise_key(key):
    text = re.sub(r"([a-z0-9])([A-Z])", r"\1 \2", key).replace("_", "-")
    return text[:1].upper() + text[1:]


def item_number(tokens, explicit_index=None):
    if explicit_index is not None:
        return explicit_index + 1
    for token in reversed(tokens):
        if isinstance(token, int):
            return token + 1
    return None


def numbered(part, number):
    if number is None:
        return part
    ar, en = part.split(" / ", 1)
    return bilingual(f"{ar} {number}", f"{en} {number}")


def part_for(filename, tokens, explicit_index=None):
    key = semantic_key(tokens)
    number = item_number(tokens, explicit_index)

    if key == "_todo":
        return bilingual("سؤال مطلوب من العميل", "Outstanding client question")

    if filename == "site.json" and tokens and tokens[0] == "ui" and key in UI_PARTS:
        return UI_PARTS[key]

    if key == "name":
        if filename == "specialties.json":
            return bilingual("اسم التخصص أو الخدمة", "Specialty or service name")
        if filename == "doctors.json":
            return bilingual("اسم الطبيب", "Doctor name")
        if filename == "branches.json":
            return bilingual("اسم الفرع", "Branch name")
        if filename == "articles.json":
            return bilingual("اسم تصنيف الموضوع", "Topic-category name")
        if filename == "digital.json":
            return bilingual("اسم المنتج أو النظام", "Product or format name")
        return bilingual("الاسم", "Name")

    if key == "title":
        if "seo" in tokens:
            return bilingual("عنوان الظهور في البحث", "Search-result title")
        if filename == "articles.json":
            return bilingual("عنوان المحتوى", "Entry title")
        if filename == "doctors.json":
            return bilingual("المسمى المهني للطبيب", "Doctor's professional title")
        return bilingual("العنوان الرئيسي", "Title")

    if key == "description" and "seo" in tokens:
        return bilingual("وصف الظهور في البحث", "Search-result description")

    if key in LIST_PARTS:
        return numbered(LIST_PARTS[key], number)

    if key == "label" and "sources" in tokens:
        return numbered(bilingual("اسم المصدر", "Source label"), number)
    if key == "label" and filename == "reviews.json":
        return numbered(bilingual("موضوع متكرر في الآراء", "Review-topic label"), number)
    if key == "label" and filename == "site.json" and "nav" in tokens:
        return bilingual("اسم رابط القائمة", "Navigation label")
    if key == "desc" and filename == "site.json" and "nav" in tokens:
        return bilingual("وصف رابط القائمة", "Navigation description")

    if key in ("q", "a"):
        return numbered(GENERIC_PARTS[key], number)
    if key in GENERIC_PARTS:
        return GENERIC_PARTS[key]

    readable = humanise_key(key)
    return bilingual(f"نص: {readable}", readable)


def aim_for(filename, tokens):
    key = semantic_key(tokens)
    if key == "_todo":
        return bilingual(
            "ده سؤال لسه محتاج إجابة مؤكدة من العيادة؛ حدّثه فقط لما تكون المعلومة الحقيقية متاحة.",
            "This records an unanswered clinic question; update it only when the confirmed fact is available.",
        )
    if key == "alt":
        return bilingual(
            "يوصل معنى الصورة لمستخدم قارئ الشاشة، أو يفضل فارغ لو الصورة للزينة فقط.",
            "It conveys a meaningful image to screen-reader users, or stays empty when the image is decorative.",
        )
    if "seo" in tokens and key == "title":
        return bilingual(
            "يقول للباحث بسرعة الصفحة عن إيه ويشجعه يفتح النتيجة من غير مبالغة.",
            "It tells a searcher what the page covers and earns the click without overclaiming.",
        )
    if "seo" in tokens and key == "description":
        return bilingual(
            "يلخّص فائدة الصفحة في نتيجة البحث قبل ما القارئ يقرر يفتحها.",
            "It previews the page's value in search results before the reader decides to visit.",
        )

    if filename == "articles.json":
        aims = {
            "name": bilingual("يسمّي التصنيف بلغة يفهمها القارئ ويجمع تحته الموضوعات المناسبة.",
                                "It names the category in reader language and groups the right topics beneath it."),
            "desc": bilingual("يوضح حدود التصنيف ونوع الإجابات اللي القارئ هيلاقيها فيه.",
                                "It sets the category's scope and the kind of answers readers will find there."),
            "title": bilingual("يمسك سؤال القارئ الحقيقي ويحدد بوضوح المحتوى هيجاوب على إيه.",
                                 "It captures the reader's real question and makes the promised answer clear."),
            "excerpt": bilingual("يدي معاينة مفيدة في الكارت تساعد القارئ يقرر هل المحتوى مناسب له.",
                                   "It gives a useful card preview so readers can judge whether the entry is relevant."),
            "heading": bilingual("يقسم الشرح لمحطات واضحة تسهّل المسح والوصول للنقطة المطلوبة.",
                                   "It divides the explanation into scannable stops so readers can find what they need."),
            "body": bilingual("يقدّم الشرح العملي الآمن اللي القارئ دخل عشانه، مع الحدود والخطوة التالية.",
                                "It delivers the practical, safe explanation readers came for, including limits and next steps."),
            "q": bilingual("يصيغ قلقاً حقيقياً بالكلمات اللي المريض غالباً بيسأل بيها.",
                             "It voices a real concern in the words a patient is likely to use."),
            "a": bilingual("يرد مباشرة وبوضوح، من غير تشخيص عن بُعد أو وعد غير مضمون.",
                             "It answers directly and clearly, without remote diagnosis or unsupported promises."),
            "tags": bilingual("يساعد في تجميع واكتشاف الموضوعات القريبة من نفس الاحتياج.",
                                "It helps related material be grouped and discovered around the same need."),
            "label": bilingual("يسمّي المرجع بوضوح عشان القارئ يقدر يعرف مصدر المعلومة.",
                                 "It identifies the reference clearly so readers can trace the information."),
        }
        if key in aims:
            return aims[key]

    if filename == "specialties.json":
        aims = {
            "name": bilingual("يسمّي التخصص أو الخدمة بنفس الصياغة في الصفحة والكروت والقائمة.",
                                "It names the specialty or service consistently across the page, cards and navigation."),
            "short": bilingual("يدي اسماً سريعاً للكروت والأماكن اللي مساحتها محدودة.",
                                 "It provides a compact name for cards and other tight spaces."),
            "sub": bilingual("يلخّص الفرق الأساسي في الخدمة تحت العنوان مباشرة.",
                               "It states the service's main distinction immediately below the title."),
            "intro": bilingual("يساعد المريض يفهم المشكلة وطريقة العيادة قبل ما يدخل في التفاصيل.",
                                 "It helps patients understand the problem and the clinic's approach before the details."),
            "treats": bilingual("يساعد القارئ يتعرف هل حالته تدخل ضمن نطاق التخصص.",
                                  "It helps readers recognise whether their concern falls within this specialty."),
            "summary": bilingual("يدي ملخصاً قابلاً للمسح يفرق الخدمة عن باقي الاختيارات.",
                                   "It gives a scannable summary that distinguishes this service from the alternatives."),
            "body": bilingual("يشرح الخدمة، حدودها، وإزاي القرار بيتحدد بعد الكشف.",
                                "It explains the service, its limits and how suitability is decided after consultation."),
            "facts": bilingual("يحوّل أهم تفاصيل الخدمة لنقاط سريعة قبل قرار الحجز.",
                                 "It turns the service's key details into quick facts before booking."),
            "q": bilingual("يعكس سؤالاً متوقعاً ممكن يمنع المريض من فهم الخدمة أو حجزها.",
                             "It surfaces a likely question that may block understanding or booking."),
            "a": bilingual("يزيل اللبس بأمان ويوضح إمتى الكشف هو الخطوة الصحيحة.",
                             "It resolves confusion safely and clarifies when a consultation is the right next step."),
        }
        if key in aims:
            return aims[key]

    if filename == "doctors.json":
        aims = {
            "name": bilingual("يعرّف الطبيب بدقة في الكروت والصفحة؛ لازم يطابق الاسم المؤكد.",
                                "It identifies the doctor accurately on cards and profiles and must match the confirmed name."),
            "title": bilingual("يلخّص الدور والتخصص المهني المؤكد عشان المريض يختار صح.",
                                 "It summarises the verified role and specialty so patients can choose appropriately."),
            "credentials": bilingual("يبني الثقة بمؤهل أو خبرة موثقة، من غير أي استنتاج أو تجميل.",
                                       "It builds trust with a verified qualification or experience point, never an inference."),
            "days": bilingual("يساعد المريض يعرف الأيام المتاحة قبل الحجز؛ لازم يطابق الجدول الحقيقي.",
                                "It shows available clinic days before booking and must match the real rota."),
            "hours": bilingual("يوضح وقت حضور الطبيب بدقة عشان ما يحصلش حجز أو زيارة غلط.",
                                 "It states the doctor's hours accurately to prevent a mistaken booking or visit."),
            "bio": bilingual("يشرح أسلوب الطبيب ونطاق خبرته بلغة إنسانية ومعلومات قابلة للإثبات.",
                               "It explains the doctor's approach and scope in human language using supportable facts."),
        }
        if key in aims:
            return aims[key]

    if filename == "branches.json":
        aims = {
            "name": bilingual("يميز الفرع بوضوح في القائمة والصفحة وخيارات الحجز.",
                                "It identifies the branch clearly in navigation, page headings and booking choices."),
            "shortName": bilingual("يوفر اسم مكان مختصر للكروت والقوائم الصغيرة.",
                                     "It provides a compact place name for cards and small menus."),
            "area": bilingual("يحدد الحي بسرعة عشان الزائر يعرف هل الموقع مناسب له.",
                                "It identifies the neighbourhood quickly so visitors can judge convenience."),
            "city": bilingual("يثبت المدينة ضمن العنوان وسياق الوصول للفرع.",
                                "It anchors the branch in the correct city for address and travel context."),
            "country": bilingual("يكمل سياق الموقع للزوار ومحركات البحث.",
                                   "It completes the location context for visitors and search engines."),
            "address": bilingual("يوصل المريض للمكان الصحيح؛ لازم يكون عنواناً مؤكداً حرفياً.",
                                   "It gets patients to the correct place and must be a confirmed exact address."),
            "landmark": bilingual("يدي علامة سهلة تساعد الزائر يلاقي المدخل من غير ارتباك.",
                                    "It gives an easy landmark that helps visitors find the entrance without confusion."),
            "hours": bilingual("يوضح وقت استقبال الفرع قبل ما المريض يتحرك أو يحجز.",
                                 "It shows when the branch receives patients before they travel or book."),
            "intro": bilingual("يعرّف دور الفرع والخدمات المتاحة أو حالة التجهيز من أول نظرة.",
                                 "It explains the branch's role, available services or opening status at a glance."),
            "gettingHere": bilingual("يدي خطوة عملية تقلل احتمالات إن الزائر يضيع في الطريق.",
                                       "It gives a practical direction that reduces the chance of a visitor getting lost."),
        }
        if key in aims:
            return aims[key]

    if filename == "digital.json":
        aims = {
            "name": bilingual("يسمّي المنتج أو نظام المتابعة بوضوح في الكارت وصفحة التفاصيل.",
                                "It names the product or programme clearly on its card and detail page."),
            "lede": bilingual("يربط العرض باحتياج حقيقي ويشرح الفائدة قبل تفاصيل المحتوى.",
                                "It connects the offer to a real need and explains the benefit before the details."),
            "short": bilingual("يلخص قيمة العرض في الكارت من غير سعر أو مبالغة.",
                                 "It summarises the offer's value on a card without price or hype."),
            "summary": bilingual("يفرق نظام المتابعة ده بسرعة عن النظام الآخر.",
                                   "It quickly distinguishes this programme format from the other option."),
            "inside": bilingual("يوضح حاجة ملموسة هياخدها المشتري داخل المنتج.",
                                  "It states one tangible thing the buyer receives inside the product."),
            "audience": bilingual("يساعد القارئ يشوف نفسه في الاستخدام المناسب للمنتج.",
                                    "It helps readers recognise whether the product suits their situation."),
            "includes": bilingual("يحدد نطاق الاشتراك بدقة عشان التوقعات تبقى واضحة قبل التواصل.",
                                    "It defines the programme scope so expectations are clear before enquiry."),
            "cta": bilingual("يقول للقارئ بوضوح إيه الخطوة التالية لو العرض مناسب له.",
                               "It tells readers the next step when the offer is relevant to them."),
            "q": bilingual("يعالج اعتراضاً أو سؤالاً شائعاً قبل طلب المنتج أو الاشتراك.",
                             "It addresses a common concern before ordering or subscribing."),
            "a": bilingual("يضبط التوقعات ويجاوب بوضوح من غير وعد طبي أو تجاري زائد.",
                             "It sets expectations and answers clearly without medical or commercial overpromising."),
        }
        if key in aims:
            return aims[key]

    if filename == "reviews.json" and key == "label":
        return bilingual("يلخص بلغة محايدة موضوعاً متكرراً فعلاً في آراء المرضى.",
                          "It neutrally names a topic that genuinely recurs in patient feedback.")

    generic = {
        "name": bilingual("يعرّف العنصر بنفس الاسم في كل مكان يظهر فيه.",
                            "It identifies the item consistently wherever it appears."),
        "shortName": bilingual("يدي نسخة مختصرة لما تكون المساحة محدودة.",
                                 "It supplies a shorter version where space is limited."),
        "kind": bilingual("يوضح بسرعة نوع المركز الطبي للزائر الجديد.",
                            "It tells a first-time visitor what kind of medical centre this is."),
        "tagline": bilingual("يلخص وعد العلامة ونبرة الرعاية في جملة سهلة التذكر.",
                               "It captures the brand promise and care philosophy in one memorable line."),
        "line1": bilingual("يحافظ على الكتابة الصحيحة للسطر الأول من الشعار النصي.",
                             "It preserves the correct wording of the wordmark's first line."),
        "line2": bilingual("يحافظ على الكتابة الصحيحة للسطر الثاني من الشعار النصي.",
                             "It preserves the correct wording of the wordmark's second line."),
        "label": bilingual("يخلي اسم الرابط أو التحكم واضحاً قبل ما القارئ يضغط عليه.",
                             "It makes a link or control understandable before the reader selects it."),
        "desc": bilingual("يدي سياقاً سريعاً يوضح للقارئ إيه اللي هيلاقيه بعد الضغط.",
                            "It previews what readers will find after following the link."),
        "description": bilingual("يشرح الغرض والقيمة بسرعة ومن غير حشو.",
                                   "It explains the purpose and value quickly, without filler."),
        "title": bilingual("يحدد موضوع الصفحة أو القسم من أول نظرة.",
                             "It establishes the page or section topic at a glance."),
        "sub": bilingual("يدعم العنوان بأهم فرق أو فائدة محتاج القارئ يعرفها.",
                           "It supports the title with the main distinction or benefit readers need."),
        "intro": bilingual("يفتح الموضوع بسياق واضح يساعد القارئ يكمل التفاصيل.",
                             "It opens with enough context for the reader to continue into the details."),
        "lede": bilingual("يشد الانتباه باحتياج حقيقي ويشرح ليه الصفحة مهمة.",
                            "It leads with a real need and explains why the page matters."),
        "short": bilingual("يدي نسخة سريعة للكروت والملخصات.",
                             "It provides a compact version for cards and previews."),
        "summary": bilingual("يوصل الفكرة الأساسية قبل التفاصيل الأطول.",
                               "It communicates the core idea before the longer detail."),
        "heading": bilingual("ينظم الصفحة ويخلي النص الطويل سهل المسح.",
                               "It organises the page and makes longer copy easy to scan."),
        "body": bilingual("يقدم المعلومة أو الخطوة العملية اللي القسم موجود عشانها.",
                            "It delivers the information or practical step the section exists to provide."),
        "q": bilingual("يصيغ السؤال بالطريقة اللي القارئ غالباً بيفكر بيها.",
                         "It frames the question in the way readers are likely to think about it."),
        "a": bilingual("يدي إجابة مباشرة ومفهومة ويوضح الخطوة التالية عند الحاجة.",
                         "It gives a direct answer and clarifies the next step where needed."),
        "bookingNote": bilingual("يضبط توقع المريض قبل الحضور ويمنع سوء فهم نظام الدخول.",
                                   "It sets expectations before arrival and prevents confusion about how visits work."),
        "blurb": bilingual("يقدم تعريفاً مختصراً بالعيادة في نهاية كل صفحة.",
                             "It gives a concise clinic introduction at the end of every page."),
        "disclaimer": bilingual("يحط حدوداً واضحة بين التوعية العامة والكشف والتشخيص الحقيقي.",
                                  "It clearly separates general education from real consultation and diagnosis."),
        "resultsVary": bilingual("يمنع توقع نتيجة مضمونة ويذكّر إن الاستجابة فردية.",
                                   "It prevents guaranteed-result expectations and reminds readers that responses differ."),
        "sampleContent": bilingual("يمنع القارئ من اعتبار المثال التوضيحي بيانات عيادة حقيقية.",
                                     "It prevents readers mistaking illustrative material for real clinic data."),
        "copyright": bilingual("يثبت اسم صاحب المحتوى وحقوقه في الفوتر.",
                                 "It identifies the content owner and rights in the footer."),
        "cta": bilingual("يوضح الخطوة المناسبة التالية من غير ضغط على القارئ.",
                           "It makes the appropriate next step clear without pressuring the reader."),
        "excerpt": bilingual("يدي معاينة قصيرة تساعد القارئ يقرر يكمل ولا لأ.",
                               "It offers a short preview so readers can decide whether to continue."),
        "question": bilingual("يجهز سؤالاً آمناً يطلب شرحاً وأسئلة للطبيب، مش تشخيصاً.",
                                "It prepares a safe prompt for explanation and doctor questions, not diagnosis."),
        "note": bilingual("يضبط توقع القارئ في اللحظة اللي يحتاج فيها تنبيه إضافي.",
                            "It sets reader expectations at the point an extra caution is needed."),
        "copied": bilingual("يؤكد فوراً إن إجراء نسخ الرابط نجح.",
                              "It immediately confirms that the link was copied successfully."),
        "credentials": bilingual("يبني الثقة بمعلومة مهنية مؤكدة وقابلة للمراجعة.",
                                   "It builds trust with confirmed, reviewable professional information."),
        "treats": bilingual("يساعد القارئ يعرف هل المحتوى أو الخدمة مرتبطة بحالته.",
                              "It helps readers judge whether the content or service relates to their concern."),
        "facts": bilingual("يعرض معلومة أساسية في صورة نقطة سريعة وواضحة.",
                             "It presents one essential detail as a quick, clear point."),
        "gettingHere": bilingual("يدي خطوة عملية تساعد الزائر يوصل بسهولة.",
                                   "It gives one practical step that helps visitors arrive easily."),
        "inside": bilingual("يوضح جزءاً ملموساً من محتوى المنتج.",
                              "It clarifies one tangible part of the product."),
        "audience": bilingual("يساعد القارئ يقرر هل العرض مناسب لاحتياجه.",
                                "It helps readers decide whether the offer fits their need."),
        "includes": bilingual("يضبط التوقعات بتحديد حاجة داخلة فعلاً في الخدمة.",
                                "It sets expectations by naming something genuinely included in the service."),
        "tags": bilingual("يساعد في وصف وتجميع المحتوى القريب من نفس الموضوع.",
                            "It helps describe and group content around the same topic."),
        "text": bilingual("ينقل الرسالة الأساسية للقارئ في المكان ده.",
                            "It communicates the essential message at this point in the page."),
    }
    if key in UI_PARTS:
        return bilingual("يخلي الإجراء أو الحالة مفهومة ومتسقة في كل صفحات الموقع.",
                          "It keeps the action or state clear and consistent throughout the site.")
    return generic.get(key, bilingual("يوصل المعلومة المقصودة بوضوح في مكان ظهورها.",
                                       "It communicates the intended information clearly where it appears."))


class Collector:
    def __init__(self, domain):
        self.domain = domain
        self.rows = []
        self.sequence = 0

    def add(self, filename, root, tokens, current_ar, current_en, ref, explicit_index=None):
        location = location_for(filename, tokens, root, self.domain)
        self.rows.append({
            "sort": (location.tour, location.entity, self.sequence),
            "page": location.heading,
            "part": part_for(filename, tokens, explicit_index),
            "link": location.link,
            "aim": aim_for(filename, tokens),
            "current_ar": scalar_text(current_ar),
            "current_en": scalar_text(current_en),
            "ref": ref,
        })
        self.sequence += 1

    def walk_todos_only(self, node, filename, root, tokens):
        """Find exact _todo keys inside an otherwise internal machine subtree."""
        if isinstance(node, dict):
            for key, value in node.items():
                child_tokens = tokens + (key,)
                if key == "_todo" and is_scalar_copy(value):
                    self.add(filename, root, child_tokens, "", value,
                             single_ref(filename, child_tokens))
                elif isinstance(value, (dict, list)):
                    self.walk_todos_only(value, filename, root, child_tokens)
        elif isinstance(node, list):
            for index, value in enumerate(node):
                if isinstance(value, (dict, list)):
                    self.walk_todos_only(value, filename, root, tokens + (index,))

    def walk(self, node, filename, root, tokens=()):
        if isinstance(node, dict):
            shape = pair_shape(node)
            if shape == "scalar":
                self.add(filename, root, tokens, node.get("ar"), node.get("en"),
                         single_ref(filename, tokens))
                for key, value in node.items():
                    if key not in ("ar", "en") and not should_skip_key(filename, tokens, key, value):
                        self.walk(value, filename, root, tokens + (key,))
                return

            if shape == "parallel-lists":
                ar_items, en_items = node["ar"], node["en"]
                common = min(len(ar_items), len(en_items))
                for index in range(common):
                    ar_tokens = tokens + ("ar", index)
                    en_tokens = tokens + ("en", index)
                    self.add(filename, root, tokens, ar_items[index], en_items[index],
                             paired_ref(filename, ar_tokens, en_tokens), explicit_index=index)
                # Preserve unmatched existing entries without inventing a mate.
                for index in range(common, len(ar_items)):
                    item_tokens = tokens + ("ar", index)
                    self.add(filename, root, tokens, ar_items[index], "",
                             single_ref(filename, item_tokens), explicit_index=index)
                for index in range(common, len(en_items)):
                    item_tokens = tokens + ("en", index)
                    self.add(filename, root, tokens, "", en_items[index],
                             single_ref(filename, item_tokens), explicit_index=index)
                for key, value in node.items():
                    if key not in ("ar", "en") and not should_skip_key(filename, tokens, key, value):
                        self.walk(value, filename, root, tokens + (key,))
                return

            for key, value in node.items():
                if should_skip_key(filename, tokens, key, value):
                    if key.startswith("_") and key != "_todo" and isinstance(value, (dict, list)):
                        self.walk_todos_only(value, filename, root, tokens + (key,))
                    continue
                if key == "_todo":
                    if is_scalar_copy(value):
                        self.add(filename, root, tokens + (key,), "", value,
                                 single_ref(filename, tokens + (key,)))
                    continue
                self.walk(value, filename, root, tokens + (key,))
            return

        if isinstance(node, list):
            for index, value in enumerate(node):
                self.walk(value, filename, root, tokens + (index,))
            return

        if isinstance(node, str) and not looks_like_machine_value(node):
            language = None
            for token in reversed(tokens):
                if token in ("ar", "en"):
                    language = token
                    break
            current_ar = node if language == "ar" else ""
            current_en = node if language != "ar" else ""
            self.add(filename, root, tokens, current_ar, current_en,
                     single_ref(filename, tokens))


def style_sheet(ws):
    widths = (42, 31, 45, 72, 68, 68, 68, 68, 32, 18)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{ws.max_row}"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 32

    for column in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BOX

    # The input headers carry the same champagne cue as their columns.
    for column in (7, 8):
        ws.cell(row=1, column=column).fill = PatternFill("solid", fgColor=CHAMPAGNE)
        ws.cell(row=1, column=column).font = Font(name="Calibri", size=11, bold=True, color=OLIVE)

    for row in range(2, ws.max_row + 1):
        fill = ALT_FILL if row % 2 == 0 else PatternFill(fill_type=None)
        for column in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row, column=column)
            cell.font = REF_FONT if column == 10 else BODY_FONT
            cell.alignment = WRAP_RTL if column in (5, 7) else WRAP
            cell.border = BOX
            cell.fill = fill
        for column in (7, 8):
            ws.cell(row=row, column=column).fill = NEW_FILL

        link_cell = ws.cell(row=row, column=3)
        if link_cell.value:
            link_cell.hyperlink = link_cell.value
            link_cell.style = "Hyperlink"
            link_cell.alignment = WRAP
            link_cell.border = BOX


def main():
    if Workbook is None:
        print("ERROR: openpyxl is required. Install it with: python -m pip install openpyxl")
        return 1

    data = {filename: load_json(filename) for filename in CONTENT_FILES}
    domain = data["site.json"].get("brand", {}).get("domain", "")
    collector = Collector(domain)
    for filename in CONTENT_FILES:
        collector.walk(data[filename], filename, data[filename])

    collector.rows.sort(key=lambda row: row["sort"])

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Content map"
    worksheet.append(list(HEADERS))
    for row in collector.rows:
        worksheet.append([
            row["page"], row["part"], row["link"], row["aim"],
            row["current_ar"], row["current_en"], "", "", "", row["ref"],
        ])
    style_sheet(worksheet)

    os.makedirs(os.path.dirname(OUTPUT), exist_ok=True)
    workbook.save(OUTPUT)
    print(f"wrote {OUTPUT}")
    print(f"editable rows: {len(collector.rows)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
