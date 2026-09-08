#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Apply a returned CONTENT-MAP.xlsx to the clinic's JSON content files.

The spreadsheet is intentionally friendly; this script is intentionally
strict.  It follows only the hidden JSON Pointer references created by
tools/export_content_map.py, refuses prices, protects copy changed after the
workbook was exported, and backs up every changed file before the first write.

Run:  python tools/apply_content_map.py
      python tools/apply_content_map.py path/to/returned.xlsx --dry-run
"""

import argparse
import collections
import datetime
import io
import json
import os
import re
import shutil
import sys
import tempfile
from dataclasses import dataclass

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT_DIR = os.path.join(ROOT, "content")
BACKUP_DIR = os.path.join(CONTENT_DIR, "_backups")
DEFAULT_WORKBOOK = os.path.join(ROOT, "_project", "CONTENT-MAP.xlsx")

CONTENT_FILES = (
    "site.json",
    "pages.json",
    "specialties.json",
    "doctors.json",
    "branches.json",
    "articles.json",
    "digital.json",
    "reviews.json",
)
ALLOWED_FILES = set(CONTENT_FILES)

HEADERS = (
    "Page / Section",
    "Part",
    "Link",
    "Aim",
    "Current (AR)",
    "Current (EN)",
    "New (AR)",
    "New (EN)",
    "Notes",
    "_ref",
)

# A price is a Latin, Arabic-Indic or Eastern Arabic digit followed by one of
# the site's forbidden public currency labels.  Separators may appear inside
# the number, but letters may not, which keeps the expression from reaching
# across an ordinary sentence to a later currency abbreviation.
DIGIT = "0-9٠-٩۰-۹"
PRICE_RE = re.compile(
    rf"[{DIGIT}](?:[{DIGIT},.٬٫]|[ \t\u00a0])*"
    rf"(?:ج[ \t\u00a0]*\.[ \t\u00a0]*م\.?|جنيه(?:اً|ا)?|EGP|L[ \t\u00a0]*\.?[ \t\u00a0]*E\.?)"
    rf"(?![A-Za-zء-ي])",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class Target:
    filename: str
    tokens: tuple
    pointer: str


class ReferenceError(ValueError):
    pass


class Summary:
    def __init__(self):
        self.rows_examined = 0
        self.rows_applied = 0
        self.skipped = collections.Counter()
        self.details = []
        self.files_changed = set()
        self.backups = []
        self.fatal = ""

    def skip(self, row_number, reason, detail="", report=False):
        self.skipped[reason] += 1
        if report or detail:
            message = f"row {row_number}: {reason}"
            if detail:
                message += f" - {detail}"
            self.details.append(message)

    @property
    def rows_skipped(self):
        return sum(self.skipped.values())

    def print(self, dry_run):
        mode = "DRY RUN" if dry_run else "WRITE"
        print(f"\nCONTENT-MAP apply summary ({mode})")
        print(f"rows examined: {self.rows_examined}")
        suffix = " (would be applied; no files written)" if dry_run else ""
        print(f"rows applied:  {self.rows_applied}{suffix}")
        print(f"rows skipped:  {self.rows_skipped}")
        for reason, count in self.skipped.items():
            print(f"  {count:>5}  {reason}")

        label = "files that would change" if dry_run else "files changed"
        changed = [name for name in CONTENT_FILES if name in self.files_changed]
        print(f"{label}: {', '.join(changed) if changed else 'none'}")
        if self.backups:
            print("backups created:")
            for path in self.backups:
                print(f"  {path}")
        if self.details:
            print("reported rows:")
            for detail in self.details:
                print(f"  {detail}")
        if self.fatal:
            print(f"ERROR: {self.fatal}")


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbook",
        nargs="?",
        default=DEFAULT_WORKBOOK,
        help="returned CONTENT-MAP.xlsx (default: _project/CONTENT-MAP.xlsx)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="report valid changes without backing up or writing JSON",
    )
    return parser.parse_args(argv)


def workbook_path(value):
    if os.path.isabs(value):
        return value
    return os.path.join(ROOT, value)


def cell_text(value, strip=True):
    if value is None:
        return ""
    text = str(value)
    return text.strip() if strip else text


def comparable(value):
    """Compare workbook and JSON text without treating edge whitespace as copy."""
    if value is None:
        return ""
    return str(value).strip()


def decode_pointer(pointer):
    if pointer == "":
        return ()
    if not pointer.startswith("/"):
        raise ReferenceError(f"invalid JSON Pointer {pointer!r}")
    tokens = []
    for raw in pointer[1:].split("/"):
        # RFC 6901 permits only ~0 and ~1 escape sequences.
        if re.search(r"~(?![01])", raw):
            raise ReferenceError(f"invalid JSON Pointer escape in {pointer!r}")
        tokens.append(raw.replace("~1", "/").replace("~0", "~"))
    return tuple(tokens)


def parse_target(text, default_filename=None):
    text = text.strip()
    if "#" not in text:
        if default_filename and text.startswith("/"):
            filename, pointer = default_filename, text
        else:
            raise ReferenceError("reference must contain a filename and # JSON Pointer")
    else:
        filename, pointer = text.split("#", 1)
        filename = filename.strip() or default_filename
    if not filename or filename not in ALLOWED_FILES:
        raise ReferenceError(f"file {filename!r} is not one of the eight content files")
    tokens = decode_pointer(pointer.strip())
    if not tokens:
        raise ReferenceError("a reference may not replace a whole JSON document")
    return Target(filename, tokens, pointer.strip())


def parse_ref(value):
    raw = cell_text(value)
    if not raw:
        raise ReferenceError("missing _ref")
    parts = [part.strip() for part in raw.split("||")]
    if len(parts) not in (1, 2) or any(not part for part in parts):
        raise ReferenceError("_ref must contain one pointer or an AR/EN pointer pair")
    first = parse_target(parts[0])
    if len(parts) == 1:
        return (first,)
    second = parse_target(parts[1], first.filename)
    if first.filename != second.filename:
        raise ReferenceError("one row may not write to two different files")
    return first, second


def list_index(token, length):
    if not re.fullmatch(r"0|[1-9][0-9]*", token):
        raise ReferenceError(f"{token!r} is not a valid array index")
    index = int(token)
    if index >= length:
        raise ReferenceError(f"array index {index} no longer exists")
    return index


def resolve(data, target):
    node = data
    for token in target.tokens:
        if isinstance(node, dict):
            if token not in node:
                raise ReferenceError(f"{target.filename}#{target.pointer} no longer resolves")
            node = node[token]
        elif isinstance(node, list):
            node = node[list_index(token, len(node))]
        else:
            raise ReferenceError(f"{target.filename}#{target.pointer} no longer resolves")
    return node


def set_existing(data, target, value):
    parent_tokens = target.tokens[:-1]
    parent_target = Target(target.filename, parent_tokens, "")
    parent = data if not parent_tokens else resolve(data, parent_target)
    final = target.tokens[-1]
    if isinstance(parent, dict):
        if final not in parent:
            raise ReferenceError(f"{target.filename}#{target.pointer} no longer resolves")
        parent[final] = value
        return
    if isinstance(parent, list):
        parent[list_index(final, len(parent))] = value
        return
    raise ReferenceError(f"{target.filename}#{target.pointer} no longer resolves")


def pointer_language(target):
    for token in reversed(target.tokens):
        if token in ("ar", "en"):
            return token
    return None


def load_content():
    data = {}
    for filename in CONTENT_FILES:
        path = os.path.join(CONTENT_DIR, filename)
        with io.open(path, "r", encoding="utf-8") as handle:
            data[filename] = json.load(handle)
    return data


def find_content_sheet(workbook):
    for worksheet in workbook.worksheets:
        values = tuple(worksheet.cell(row=1, column=index).value for index in range(1, 11))
        if values == HEADERS:
            return worksheet
    expected = ", ".join(HEADERS)
    raise ValueError(f"no worksheet has the required header row: {expected}")


def is_formula(cell):
    return cell.data_type == "f" and cell.value not in (None, "")


def ensure_copy_scalar(value, target):
    if value is not None and not isinstance(value, str):
        raise ReferenceError(
            f"{target.filename}#{target.pointer} is no longer a text field"
        )


def make_assignment(target, live_value, expected, new_value, language):
    ensure_copy_scalar(live_value, target)
    if comparable(live_value) != comparable(expected):
        raise ReferenceError(
            f"live {language.upper()} text differs from Current; the workbook is stale"
        )
    if comparable(live_value) == comparable(new_value):
        return None
    return target, new_value, language


def row_assignments(targets, content, current_ar, current_en, new_ar, new_en):
    """Resolve a row to exact scalar assignments without mutating content."""
    candidates = {}
    if new_ar and comparable(new_ar) != comparable(current_ar):
        candidates["ar"] = new_ar
    if new_en and comparable(new_en) != comparable(current_en):
        candidates["en"] = new_en
    if not candidates:
        return [], "unchanged from Current"

    # Resolve every component of _ref before considering a partial-language
    # update.  A damaged reference is reported as a whole row and never guessed.
    resolved = [(target, resolve(content[target.filename], target)) for target in targets]

    if len(targets) == 2:
        by_language = {}
        for target, value in resolved:
            language = pointer_language(target)
            if language not in ("ar", "en") or language in by_language:
                raise ReferenceError("paired _ref does not identify one AR and one EN field")
            ensure_copy_scalar(value, target)
            by_language[language] = (target, value)
        if set(by_language) != {"ar", "en"}:
            raise ReferenceError("paired _ref does not identify one AR and one EN field")

        assignments = []
        currents = {"ar": current_ar, "en": current_en}
        for language in ("ar", "en"):
            if language not in candidates:
                continue
            target, live_value = by_language[language]
            assignment = make_assignment(
                target, live_value, currents[language], candidates[language], language
            )
            if assignment:
                assignments.append(assignment)
        return assignments, "already matches live content" if not assignments else ""

    target, node = resolved[0]
    if isinstance(node, dict):
        if "ar" not in node or "en" not in node:
            raise ReferenceError("referenced object is no longer a bilingual {ar, en} field")
        assignments = []
        currents = {"ar": current_ar, "en": current_en}
        for language in ("ar", "en"):
            if language not in candidates:
                continue
            child_target = Target(
                target.filename,
                target.tokens + (language,),
                target.pointer.rstrip("/") + "/" + language,
            )
            live_value = node[language]
            assignment = make_assignment(
                child_target, live_value, currents[language], candidates[language], language
            )
            if assignment:
                assignments.append(assignment)
        return assignments, "already matches live content" if not assignments else ""

    language = pointer_language(target)
    if language in ("ar", "en"):
        other = "en" if language == "ar" else "ar"
        if other in candidates:
            raise ReferenceError(
                f"New ({other.upper()}) cannot target this {language.upper()}-only field"
            )
        if language not in candidates:
            return [], "new text was entered in the wrong language column"
        current = current_ar if language == "ar" else current_en
        assignment = make_assignment(
            target, node, current, candidates[language], language
        )
        return ([assignment] if assignment else []), (
            "" if assignment else "already matches live content"
        )

    # A genuinely shared/monolingual scalar is shown in Current (EN) by the
    # exporter.  Accept either New column for convenience, but never guess
    # between two different proposed values.
    if "ar" in candidates and "en" in candidates:
        if comparable(candidates["ar"]) != comparable(candidates["en"]):
            raise ReferenceError(
                "shared text has different AR and EN proposals; enter one value only"
            )
        proposed = candidates["ar"]
        proposed_language = "shared"
    elif "ar" in candidates:
        proposed = candidates["ar"]
        proposed_language = "shared"
    else:
        proposed = candidates["en"]
        proposed_language = "shared"

    nonempty_currents = [value for value in (current_ar, current_en) if comparable(value)]
    if len(nonempty_currents) == 2 and comparable(nonempty_currents[0]) != comparable(nonempty_currents[1]):
        raise ReferenceError("shared field has two different Current values")
    expected = nonempty_currents[0] if nonempty_currents else ""
    assignment = make_assignment(target, node, expected, proposed, proposed_language)
    return ([assignment] if assignment else []), (
        "" if assignment else "already matches live content"
    )


def backup_changed_files(filenames):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y-%m-%dT%H-%M-%S-%f")
    created = []
    # Complete every backup before allowing the first content write.
    for filename in filenames:
        source = os.path.join(CONTENT_DIR, filename)
        destination = os.path.join(BACKUP_DIR, f"{filename}.{stamp}.bak")
        shutil.copy2(source, destination)
        created.append(destination)
    return created


def write_json_atomic(filename, data):
    """Write UTF-8, indent=1 JSON with one trailing LF, then replace atomically."""
    destination = os.path.join(CONTENT_DIR, filename)
    descriptor, temporary = tempfile.mkstemp(
        prefix=f".{filename}.", suffix=".tmp", dir=CONTENT_DIR
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as handle:
            json.dump(data, handle, ensure_ascii=False, indent=1)
            handle.write("\n")
        os.replace(temporary, destination)
    except Exception:
        try:
            os.unlink(temporary)
        except OSError:
            pass
        raise


def apply_rows(worksheet, content, summary):
    for row_number in range(2, worksheet.max_row + 1):
        cells = [worksheet.cell(row=row_number, column=column) for column in range(1, 11)]
        if not any(cell.value not in (None, "") for cell in cells):
            continue
        summary.rows_examined += 1

        current_ar = cell_text(cells[4].value, strip=False)
        current_en = cell_text(cells[5].value, strip=False)
        new_ar = cell_text(cells[6].value)
        new_en = cell_text(cells[7].value)
        ref_value = cells[9].value

        if is_formula(cells[6]) or is_formula(cells[7]):
            summary.skip(row_number, "formula in a New cell", "replace the formula with plain text", True)
            continue
        if not new_ar and not new_en:
            summary.skip(row_number, "both New cells empty")
            continue

        changed_proposals = []
        if new_ar and comparable(new_ar) != comparable(current_ar):
            changed_proposals.append(("AR", new_ar))
        if new_en and comparable(new_en) != comparable(current_en):
            changed_proposals.append(("EN", new_en))
        if not changed_proposals:
            summary.skip(row_number, "unchanged from Current")
            continue

        price_languages = [language for language, text in changed_proposals if PRICE_RE.search(text)]
        if price_languages:
            summary.skip(
                row_number,
                "price-like text refused",
                f"currency pattern found in New ({', '.join(price_languages)})",
                True,
            )
            continue

        try:
            targets = parse_ref(ref_value)
            assignments, no_change_reason = row_assignments(
                targets, content, current_ar, current_en, new_ar, new_en
            )
        except ReferenceError as error:
            message = str(error)
            reason = "stale workbook" if "workbook is stale" in message else "unresolved or unsafe _ref"
            summary.skip(row_number, reason, message, True)
            continue

        if not assignments:
            summary.skip(row_number, no_change_reason or "no applicable change")
            continue

        for target, value, _language in assignments:
            set_existing(content[target.filename], target, value)
            summary.files_changed.add(target.filename)
        summary.rows_applied += 1


def run(args, summary):
    if load_workbook is None:
        raise RuntimeError(
            "openpyxl is required. Install it with: python -m pip install openpyxl"
        )

    path = workbook_path(args.workbook)
    workbook = load_workbook(path, data_only=False)
    worksheet = find_content_sheet(workbook)
    content = load_content()
    apply_rows(worksheet, content, summary)

    changed = [filename for filename in CONTENT_FILES if filename in summary.files_changed]
    if args.dry_run or not changed:
        return

    summary.backups = backup_changed_files(changed)
    for filename in changed:
        write_json_atomic(filename, content[filename])


def main(argv=None):
    args = parse_args(argv)
    summary = Summary()
    exit_code = 0
    try:
        run(args, summary)
    except Exception as error:
        summary.fatal = str(error)
        exit_code = 1
    finally:
        # Even dependency, workbook and write failures produce the requested
        # summary instead of exiting without an explanation.
        summary.print(args.dry_run)
    return exit_code


if __name__ == "__main__":
    sys.exit(main())
